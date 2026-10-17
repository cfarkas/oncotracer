#!/usr/bin/env python3
"""Reconcile exact-allele calls with imputation; infer evidence, never confirmed origin."""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import heapq
import html
import itertools
import json
import math
from pathlib import Path
import re
import subprocess
from urllib.parse import quote

import pipeline as p
import traits_lpwgs as traits

VERSION = 2
CATEGORIES = ('germline_supported', 'somatic_candidate', 'unconfirmed')
FIELDS = ['chrom','pos','ref','alt','category','evidence_tier','reason','gene','consequence',
          'clinvar_classification','clinvar_review','review_stars','condition','clinvar_allele_id','clinvar_url',
          'germline_GT','germline_DP','germline_GQ','germline_AD','germline_filter',
          'somatic_GT','somatic_DP','somatic_AD','somatic_VAF','somatic_filter',
          'imputation_state','imputed_GT','imputed_GP','imputed_DS','imputed_INFO',
          'called_by','pathogenic_label','clinical_confirmation','explanation']
SOURCES = [
    'https://odelaneau.github.io/GLIMPSE/docs/tutorials/getting_started/',
    'https://gatk.broadinstitute.org/hc/en-us/articles/360037593851-Mutect2',
    'https://www.ncbi.nlm.nih.gov/clinvar/docs/review_status/',
]


def annotation_rows(path):
    with open(path, newline='') as handle:
        reader=csv.DictReader(handle, delimiter='\t')
        required={'chrom','pos','ref','alt','genotype','DP','GQ','AD','filter','genotype_status'}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f'Incomplete annotation columns: {path}')
        for row in reader:
            for key in ['ref','alt','genotype']:row[key]=row[key].upper()
            yield row


def locus(row, order):
    return order[row['chrom']], int(row['pos'])


def groups(rows, kind, order):
    previous=None
    for key, values in itertools.groupby(rows, key=lambda r:locus(r,order)):
        if previous is not None and key<previous:
            raise ValueError(f'{kind} input is not sorted by reference contig and position')
        previous=key
        yield key, kind, list(values)


def joined(streams):
    for key, group in itertools.groupby(heapq.merge(*streams, key=lambda item:item[0]), key=lambda item:item[0]):
        result=defaultdict(list)
        for _, kind, rows in group: result[kind].extend(rows)
        yield key, result


def fraction(row):
    try:
        ad=[float(v) for v in row.get('AD','').split(',')]
        if len(ad)!=2 or any(not math.isfinite(x) or x<0 for x in ad) or sum(ad)==0: return None
        return ad[1]/sum(ad)
    except (ValueError,TypeError): return None


def observed_alt(row):
    return bool(row and row['alt'] in re.split(r'[/|]',row.get('genotype','.')))


def read_germline_supported(row,c):
    if not row or row.get('genotype_status')!='called' or row.get('filter')!='PASS' or not observed_alt(row): return False
    if p.num(row.get('DP'))<c['min_depth'] or p.num(row.get('GQ'))<c['min_gq']: return False
    af=fraction(row)
    if af is None: return False
    alleles=re.split(r'[/|]',row['genotype'])
    if len(alleles)!=2 or not set(alleles)<={row['ref'],row['alt']}: return False
    return .25<=af<=.75 if len(set(alleles))==2 else af>=.8


def somatic_supported(row,c):
    if not row or row.get('filter')!='PASS' or not observed_alt(row): return False
    try:
        ad=[float(x) for x in row.get('AD','').split(',')]
        return p.num(row.get('DP'))>=c['min_depth'] and len(ad)==2 and all(math.isfinite(x) and x>=0 for x in ad) and ad[1]>=3
    except (ValueError,TypeError): return False


def imputation_evidence(ref,alt,records,c):
    empty={'state':'not_in_imputed_output','GT':'.','GP':'.','DS':'.','INFO':'.','alleles':[]}
    if len(ref)!=1 or len(alt)!=1 or ref not in 'ACGT' or alt not in 'ACGT':
        return {**empty,'state':'not_assessed_non_snp'}
    variants={}
    for rec in records: traits.add_imputed_record(variants,rec)
    if not records: return empty
    by_allele=variants[(records[0]['chrom'],records[0]['pos'])]
    if len(by_allele)>1 or any(v['conflict'] for v in by_allele.values()):
        return {**empty,'state':'ambiguous_site'}
    entry=by_allele.get((ref,alt))
    if entry is None: return {**empty,'state':'allele_mismatch'}
    call=traits.imputed_call(entry['record'],c)
    result={**empty,'state':'low_confidence','GT':call['genotype'],'GP':call['GP'],'DS':call['DS'],'INFO':call['INFO'],'alleles':call['alleles']}
    if call['status']=='imputed':
        result['state']='alternate_supported' if alt in call['alleles'] else 'reference_supported'
    return result


def has_clinvar(row):
    # ClinVar includes records with an AlleleID but no germline CLNSIG assertion.
    return any(row.get(key) not in {None,'','.'} for key in ('clinvar_allele_id','clinvar_classification'))


def pathogenic_label(value):
    return any(x.strip().lower().replace(' ','_') in {'pathogenic','likely_pathogenic'} for x in re.split(r'[/|,;]',value))


def review_stars(value):
    value=value.lower()
    if any(x in value for x in ('no_assertion_criteria_provided','no_assertion_provided','no_classification_provided')): return 0
    return p.review_stars(value)


def decide(g,s,imp,c,duplicate=False):
    if duplicate: return 'unconfirmed','conflicting_records','duplicate_caller_records'
    good_g=read_germline_supported(g,c);good_s=somatic_supported(s,c)
    if good_g and imp['state'] in {'alternate_supported','reference_supported'} and sorted(re.split(r'[/|]',g['genotype']))!=sorted(imp['alleles']):
        return 'unconfirmed','discordant_evidence','germline_reads_disagree_with_imputation'
    if good_g:
        return 'germline_supported','reads_and_imputation' if imp['state']=='alternate_supported' else 'direct_reads','diploid_germline_call_passes_read_filters'
    if imp['state']=='alternate_supported':
        return 'germline_supported','imputed_support','exact_allele_has_passing_imputed_nonreference_genotype'
    if good_s and imp['state'] not in {'ambiguous_site','allele_mismatch'}:
        return 'somatic_candidate','unpaired_read_support','mutect_pass_and_read_support_without_germline_support'
    if imp['state']=='ambiguous_site': return 'unconfirmed','ambiguous_imputation','multiple_or_conflicting_imputed_alleles_at_position'
    return 'unconfirmed','insufficient_evidence','read_or_imputation_support_does_not_resolve_origin'


def unique_record(rows):
    if not rows: return None,False
    return rows[0], any(r!=rows[0] for r in rows[1:])


def reconcile(g_rows,s_rows,imputed,c):
    g,g_conflict=unique_record(g_rows);s,s_conflict=unique_record(s_rows)
    base=g or s;ref,alt=base['ref'],base['alt']
    imp=imputation_evidence(ref,alt,imputed,c)
    category,tier,reason=decide(g,s,imp,c,g_conflict or s_conflict)
    def annotation(key):
        vals=[]
        for source in [g,s]:
            val=(source or {}).get(key,'.')
            if val not in {'','.'} and val not in vals: vals.append(val)
        return ' | '.join(vals) or '.'
    row={k:base[k] for k in ['chrom','pos','ref','alt']}
    row.update(category=category,evidence_tier=tier,reason=reason,called_by='+'.join(k for k,r in [('germline',g),('somatic',s)] if r))
    for key in ['gene','consequence','clinvar_classification','clinvar_review','condition','clinvar_allele_id']:
        row[key]=annotation(key).replace('\\x3b',';')
    row['review_stars']=review_stars(row['clinvar_review'])
    allele_ids=re.findall(r'\d+',row['clinvar_allele_id'])
    row['clinvar_url']='https://www.ncbi.nlm.nih.gov/clinvar/?term='+quote(' OR '.join(x+'[AlleleID]' for x in allele_ids),safe='') if allele_ids else '.'
    for prefix,record in [('germline',g),('somatic',s)]:
        for name in ['GT','DP','GQ','AD','filter']:
            if prefix=='somatic' and name=='GQ': continue
            row[prefix+'_'+name]=(record or {}).get('genotype' if name=='GT' else name,'.')
    af=fraction(s or {});row['somatic_VAF']=round(af,6) if af is not None else '.'
    row['imputation_state']=imp['state']
    row.update({'imputed_'+k:imp[k] for k in ['GT','GP','DS','INFO']})
    row['pathogenic_label']='yes' if pathogenic_label(row['clinvar_classification']) else 'no'
    row['clinical_confirmation']='required' if row['pathogenic_label']=='yes' else 'not_clinically_validated'
    meanings={
        'germline_supported':'Evidence is compatible with a germline genotype; inherited origin is inferred, not proven.',
        'somatic_candidate':'Unpaired Mutect2 evidence passes the stated read rules; somatic origin is unconfirmed and a germline variant or artifact can remain.',
        'unconfirmed':'Available evidence is insufficient, conflicting, or ambiguous; the variant is retained for review.'}
    row['explanation']=meanings[category]
    if tier=='imputed_support': row['explanation']+=' Support comes from imputation; the sequence-read call may remain low depth or filtered. Imputation uses the same reads and is not independent confirmation.'
    if row['pathogenic_label']=='yes': row['explanation']+=' ClinVar contains a pathogenic/likely pathogenic term. Confirm the genotype and review the exact condition, inheritance, and database evidence before any clinical interpretation.'
    if 'conflict' in row['clinvar_classification'].lower(): row['explanation']+=' ClinVar classifications conflict.'
    return row


def prepare_imputed(a,c,order):
    bed=Path('variant_sites.bed')
    with bed.open('w') as handle:
        streams=[groups(annotation_rows(a.germline),'g',order),groups(annotation_rows(a.somatic),'s',order)]
        for (_,position),rows in joined(streams):
            row=next(iter(rows.values()))[0]
            handle.write(f'{row["chrom"]}\t{position-1}\t{position}\n')
    out=Path('imputation_at_called_sites.tsv.gz')
    bc=p.tool(c,'bcftools')
    # Query only observed-call positions, with exact REF/ALT matching performed below.
    by_chrom={}
    for file in a.imputed:
        match=re.search(r'\.(chr\d+|\d+)\.imputed\.bcf$',file)
        if not match: raise ValueError(f'Cannot identify imputed chromosome: {file}')
        chrom=match.group(1)
        if chrom in by_chrom: raise ValueError(f'Duplicate chromosome file: {chrom}')
        samples=p.capture([bc,'query','-l',file]).splitlines()
        if a.sample not in samples: raise ValueError(f'Wrong sample in imputed file: {file}')
        by_chrom[chrom]=file
    if c.get('glimpse2') and set(by_chrom)!=set(c['glimpse_chromosomes']):
        raise ValueError('Imputation chromosome set is incomplete; do not classify a partial run')
    with gzip.open(out,'wt',compresslevel=1) as handle:
        if bed.stat().st_size:
            for chrom,file in sorted(by_chrom.items(),key=lambda kv:order[kv[0]]):
                print('Reconciling imputation:',chrom,flush=True)
                command=[bc,'query','--allow-undef-tags','-s',a.sample,'-R',str(bed),'--regions-overlap','0','-f','%CHROM\t%POS\t%REF\t%ALT\t%FILTER\t%INFO/INFO[\t%GT\t%DS\t%GP]\n',file]
                with subprocess.Popen(command,stdout=subprocess.PIPE,text=True) as proc:
                    for line in proc.stdout: handle.write(line)
                    if proc.wait(): raise subprocess.CalledProcessError(proc.returncode,command)
    return out


def imputed_rows(path):
    with gzip.open(path,'rt') as handle:
        for line in handle:
            chrom,pos,ref,alt,filt,info,gt,ds,gp=line.rstrip('\n').split('\t')
            yield {'chrom':chrom,'pos':int(pos),'ref':ref,'alt':alt,'filter':filt,
                   'info':{} if info=='.' else {'INFO':info},'fmt':{'GT':gt,'DS':ds,'GP':gp}}


class Workbook:
    NUMERIC_FIELDS={'pos','review_stars','germline_DP','germline_GQ','somatic_DP','somatic_VAF','imputed_GP','imputed_DS','imputed_INFO'}
    def __init__(self,path,sample):
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.cell import WriteOnlyCell
        self.cell=WriteOnlyCell;self.path=path;self.book=Workbook(write_only=True);self.counts=Counter();self.sheets={}
        self.font=Font(bold=True,color='FFFFFF');self.fill=PatternFill('solid',fgColor='173F56')
        guide=self.book.create_sheet('Read_me')
        guide.column_dimensions['A'].width=26;guide.column_dimensions['B'].width=115
        for row in [
            [sample+' variant evidence review'],
            ['Scope','All ClinVar-matched exact alleles from the germline and somatic call union; full genome-wide category tables are compressed TSV files.'],
            ['Germline_supported','Direct diploid read support or passing exact-allele GLIMPSE2 support. Not confirmed inherited origin.'],
            ['Somatic_candidates','Unpaired Mutect2 PASS with DP at least the configured minimum and at least 3 ALT reads, without supporting germline evidence. Not confirmed somatic origin.'],
            ['Unconfirmed','Low evidence, ambiguous imputation, or contradictory data. Missing imputation is not evidence of somatic origin.'],
            ['Pathogenic_LP','All explicit ClinVar pathogenic/likely pathogenic terms, including low-evidence records. Not a list of confirmed disease findings.'],
            ['Missing ClinVar classification','An AlleleID match is retained even when the clinical-significance field is blank (.). No pathogenicity or benignity is inferred from that blank field.'],
            ['Imputation','The same sample reads contribute to imputation. Imputation is not independent genotype confirmation and assesses selected autosomal SNPs only.'],
            ['Clinical classification','Evidence category describes genotype/origin support, not whether a variant causes disease.'],
            ['Review stars','No assertion criteria provided is zero stars.'],
            *[['Method source',url] for url in SOURCES]]:guide.append(row)
        self.summary=self.book.create_sheet('Summary')
        self.summary.column_dimensions['A'].width=55
        self.summary.column_dimensions['B'].width=32
        self.summary.column_dimensions['C'].width=28
        for name in ['ClinVar_all','Germline_supported','Somatic_candidates','Unconfirmed','Pathogenic_LP']:
            self.create_sheet(name)
    def create_sheet(self,name):
        sheet=self.book.create_sheet(name);sheet.freeze_panes='E2'
        headers=[]
        for field in FIELDS:
            cell=self.cell(sheet,value=field);cell.font=self.font;cell.fill=self.fill;headers.append(cell)
        from openpyxl.utils import get_column_letter
        widths={'reason':58,'explanation':110,'condition':60,'clinvar_url':60,'clinvar_classification':38,'clinvar_review':46}
        for i,field in enumerate(FIELDS,1):sheet.column_dimensions[get_column_letter(i)].width=widths.get(field,22 if i>4 else 13)
        sheet.append(headers)
        self.sheets[name]=sheet
        return sheet
    def add(self,name,row):
        # Split before Excel's row limit instead of silently dropping rows.
        ordinal=self.counts[name]//1048575+1;sheetname=name if ordinal==1 else f'{name[:26]}_{ordinal}'
        sheet=self.sheets.get(sheetname) or self.create_sheet(sheetname)
        values=[]
        for field in FIELDS:
            value=row.get(field,'.')
            if field in self.NUMERIC_FIELDS and isinstance(value,str) and value not in {'','.'}:
                try:
                    number=float(value)
                    if math.isfinite(number):value=int(number) if number.is_integer() else number
                except ValueError:pass
            if isinstance(value,str):
                cell=self.cell(sheet,value=value[:32767]);cell.data_type='s';values.append(cell)
            else:values.append(value)
        sheet.append(values);self.counts[name]+=1
    def close(self,summary):
        from openpyxl.utils import get_column_letter
        self.summary.append(['Category','All unique observed variants','ClinVar-matched variants'])
        for category in CATEGORIES:self.summary.append([category,summary['category_counts'].get(category,0),summary['clinvar_by_category'].get(category,0)])
        self.summary.append(['All unique variants',summary['unique_called_variants']])
        self.summary.append(['Pathogenic/likely pathogenic labels (not confirmed)',summary['pathogenic_label_variants']])
        for name,sheet in self.sheets.items():
            match=re.search(r'_(\d+)$',name);base=name[:match.start()] if match else name;ordinal=int(match.group(1)) if match else 1
            n=min(1048575,max(0,self.counts[base]-(ordinal-1)*1048575))
            sheet.auto_filter.ref=f'A1:{get_column_letter(len(FIELDS))}{n+1}'
        self.book.save(self.path)


def run(a,c):
    order={name:i for i,name in enumerate(p.contigs(c))}
    imputed=prepare_imputed(a,c,order)
    handles={};writers={}
    names=[*CATEGORIES,'all_variants','clinvar','pathogenic_labels']
    for name in names:
        handle=gzip.open(f'{a.sample}.{name}.tsv.gz','wt',compresslevel=1,newline='');handles[name]=handle
        writer=csv.DictWriter(handle,fieldnames=FIELDS,delimiter='\t');writer.writeheader();writers[name]=writer
    workbook=Workbook(f'{a.sample}.variant_review.xlsx',a.sample)
    counts=Counter();clinvar_counts=Counter();tiers=Counter();states=Counter();reasons=Counter();total=0;pathogenic=0;source_rows=Counter()
    streams=[groups(annotation_rows(a.germline),'g',order),groups(annotation_rows(a.somatic),'s',order),groups(imputed_rows(imputed),'i',order)]
    try:
        for _,rows in joined(streams):
            variants=defaultdict(lambda:defaultdict(list))
            for kind in ['g','s']:
                for row in rows[kind]:
                    variants[(row['ref'],row['alt'])][kind].append(row);source_rows[kind]+=1
            for allele,by_source in sorted(variants.items()):
                row=reconcile(by_source['g'],by_source['s'],rows['i'],c)
                category=row['category'];counts[category]+=1;tiers[row['evidence_tier']]+=1;states[row['imputation_state']]+=1;reasons[row['reason']]+=1;total+=1
                writers['all_variants'].writerow(row);writers[category].writerow(row)
                if has_clinvar(row):
                    clinvar_counts[category]+=1;writers['clinvar'].writerow(row);workbook.add('ClinVar_all',row)
                    workbook.add({'germline_supported':'Germline_supported','somatic_candidate':'Somatic_candidates','unconfirmed':'Unconfirmed'}[category],row)
                if row['pathogenic_label']=='yes':
                    pathogenic+=1;writers['pathogenic_labels'].writerow(row);workbook.add('Pathogenic_LP',row)
                if total%500000==0:print(f'Classified {total:,} unique called variants',flush=True)
    finally:
        for handle in handles.values():handle.close()
    summary={'version':VERSION,'sample':a.sample,'test_mode':c['test_mode'],'unique_called_variants':total,
        'category_counts':dict(counts),'clinvar_by_category':dict(clinvar_counts),'evidence_tiers':dict(tiers),
        'imputation_states':dict(states),'reasons':dict(reasons),'source_annotation_rows':dict(source_rows),
        'pathogenic_label_variants':pathogenic,'imputed_chromosomes':len(a.imputed),'confirmed_somatic_variants':0,
        'scope':'Union of normalized exact alleles present in sequence-based germline or somatic caller outputs. Imputed-only sites remain in the chromosome BCFs, not in these observed-call tables.',
        'thresholds':{k:c[k] for k in ['min_depth','min_gq','imputation_min_gp','imputation_min_info']},
        'interpretation':'Evidence categories are research inferences, not verified biological origin, disease diagnoses, or independent confirmation. No matched normal is available.',
        'sources':SOURCES}
    workbook.close(summary)
    Path(f'{a.sample}.variant_evidence.json').write_text(json.dumps(summary,indent=2)+'\n')
    Path(f'{a.sample}.variant_evidence.html').write_text('<!doctype html><html><meta charset="utf-8"><title>Variant evidence</title><body><h1>'+html.escape(a.sample)+' variant evidence</h1><p>'+html.escape(summary['interpretation'])+'</p><pre>'+html.escape(json.dumps(summary,indent=2))+'</pre></body></html>')
    print(json.dumps(summary,indent=2),flush=True)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',required=True);parser.add_argument('--sample',required=True)
    parser.add_argument('--germline',required=True);parser.add_argument('--somatic',required=True)
    parser.add_argument('--imputed',nargs='*',default=[])
    a=parser.parse_args();run(a,json.loads(Path(a.config).read_text()))

if __name__=='__main__':main()
