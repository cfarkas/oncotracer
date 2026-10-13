"""Shared validated genotype access and transparent exports for genome extensions."""
import csv
import gzip
import hashlib
import html
import json
import math
from pathlib import Path

import pysam


def sha256(path):
    digest=hashlib.sha256()
    with open(path,'rb') as handle:
        for block in iter(lambda:handle.read(4*1024*1024),b''):digest.update(block)
    return digest.hexdigest()


def runtime_signature(config):
    return hashlib.sha256(json.dumps(config,sort_keys=True,separators=(',',':')).encode()).hexdigest()


def dump(path,value):
    Path(path).write_text(json.dumps(value,indent=2,allow_nan=False)+'\n')


def flat(value):
    if value is None:return '.'
    if isinstance(value,(list,tuple)):return ','.join(map(flat,value))
    return str(value)


def record_dict(rec,sample):
    if sample not in rec.samples:raise ValueError(f'Sample {sample} absent from {rec.contig}:{rec.pos}')
    fmt=rec.samples[sample];gt=fmt.get('GT') or ()
    return {'chrom':rec.contig,'pos':rec.pos,'ref':rec.ref,'alt':','.join(rec.alts or []),
        'filter':';'.join(rec.filter.keys()) or '.',
        'info':{k:flat(v) for k,v in rec.info.items()},
        'fmt':{**{k:flat(v) for k,v in fmt.items()},'GT':('|' if fmt.phased else '/').join('.' if x is None else str(x) for x in gt)}}


def posterior(rec,sample,config):
    """Validate biallelic diploid GP/DS; never replace a missing dosage by zero."""
    from traits_lpwgs import imputed_call
    result=imputed_call(record_dict(rec,sample),config)
    fmt=dict(rec.samples[sample])
    if result['status']=='imputed':
        gp=fmt.get('GP');ds=fmt.get('DS')
        if isinstance(ds,tuple):ds=ds[0] if len(ds)==1 else None
        if gp is None or ds is None or not math.isfinite(ds) or abs(ds-(gp[1]+2*gp[2]))>.03:
            result['status']='inconsistent_or_missing_dosage'
    return result


def unique_inputs(paths,sample):
    """One imputed BCF per chromosome with its index and the expected single sample."""
    result={}
    for path in paths:
        with pysam.VariantFile(str(path)) as vcf:
            if list(vcf.header.samples)!=[sample]:raise ValueError(f'Expected only sample {sample} in {path}')
            first=next(iter(vcf),None)
            if first is None:raise ValueError(f'Empty imputed chromosome: {path}')
            chrom=first.contig
            if chrom in result:raise ValueError(f'Duplicate imputed chromosome input: {chrom}')
            if not any(Path(str(path)+ext).is_file() for ext in ['.csi','.tbi']):raise ValueError(f'Missing imputation index: {path}')
            result[chrom]=str(path)
    return result


def table(path,fields,rows):
    opener=gzip.open if str(path).endswith('.gz') else open
    with opener(path,'wt',newline='') as handle:
        writer=csv.DictWriter(handle,fieldnames=fields,delimiter='\t',extrasaction='ignore')
        writer.writeheader();writer.writerows(rows)


def page(path,title,introduction,rows,fields,links=()):
    e=lambda value:html.escape(str(value))
    cells=''.join('<tr>'+''.join('<td>'+e(row.get(k,'.'))+'</td>' for k in fields)+'</tr>' for row in rows)
    nav=' · '.join('<a href="'+e(url)+'">'+e(label)+'</a>' for label,url in links)
    Path(path).write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>'+e(title)+'</title><style>body{font:16px/1.6 system-ui;margin:2rem;color:#173044}table{border-collapse:collapse}th,td{border:1px solid #dce5ea;padding:.6rem;vertical-align:top}th{background:#eef4f7}td{max-width:45rem;overflow-wrap:anywhere}</style><h1>'+e(title)+'</h1><p>'+e(introduction)+'</p><p>'+nav+'</p><div style="overflow:auto"><table><thead><tr>'+''.join('<th>'+e(k)+'</th>' for k in fields)+'</tr></thead><tbody>'+cells+'</tbody></table></div></html>')


class Excel:
    """Streaming workbook with safe strings, filters, numeric cells and row splitting."""
    def __init__(self,path,notes):
        from openpyxl import Workbook
        self.book=Workbook(write_only=True);self.path=path;self.sheets={};self.counts={};self.fields={}
        guide=self.book.create_sheet('Read_me');guide.column_dimensions['A'].width=30;guide.column_dimensions['B'].width=120
        for row in notes:self.append_row(guide,row)
    def append_row(self,sheet,values):
        from openpyxl.cell import WriteOnlyCell
        cells=[]
        for value in values:
            if isinstance(value,str):
                cell=WriteOnlyCell(sheet,value=value[:32767]);cell.data_type='s';cells.append(cell)
            else:cells.append(value)
        sheet.append(cells)
    def create(self,name,fields):
        self.fields[name]=fields;self.counts[name]=0;self._sheet(name,fields)
    def _sheet(self,name,fields):
        from openpyxl.utils import get_column_letter
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font,PatternFill
        sheet=self.book.create_sheet(name);sheet.freeze_panes='E2';cells=[]
        for i,key in enumerate(fields,1):
            sheet.column_dimensions[get_column_letter(i)].width=60 if any(w in key for w in ['condition','interpretation','explanation','population','ancestr','reason']) else 24
            cell=WriteOnlyCell(sheet,value=key);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173F56');cells.append(cell)
        sheet.append(cells);self.sheets[name]=sheet;return sheet
    def add(self,name,row):
        number=self.counts[name]//1048575+1;actual=name if number==1 else name[:25]+'_'+str(number)
        sheet=self.sheets.get(actual) or self._sheet(actual,self.fields[name])
        self.append_row(sheet,[row.get(k,'.') for k in self.fields[name]]);self.counts[name]+=1
    def close(self):
        from openpyxl.utils import get_column_letter
        for name,fields in self.fields.items():
            for number in range(1,self.counts[name]//1048575+2):
                actual=name if number==1 else name[:25]+'_'+str(number)
                if actual in self.sheets:
                    n=min(1048575,max(0,self.counts[name]-(number-1)*1048575))
                    self.sheets[actual].auto_filter.ref=f'A1:{get_column_letter(len(fields))}{n+1}'
        self.book.save(self.path)
