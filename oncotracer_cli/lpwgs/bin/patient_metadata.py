"""Optional patient context; never infer names or ages from genetic data."""
import csv
import datetime
import hashlib
import re
from pathlib import Path

FIELDS=['sample_id','name','age','sex','birth_year','age_context','age_as_of','source']
MISSING={'','none','nan','missing','no_aplica','no aplica','unknown','na','n/a','.'}

def text(v):return '' if v is None or str(v).strip().lower() in MISSING else str(v).strip()
def canonical(v):
 value=re.sub(r'[-_\s]','',text(v)).upper();m=re.fullmatch(r'([A-Z]+)0*(\d+)',value)
 return m[1]+f'{int(m[2]):03d}' if m else value

def sex(v):
 v=text(v).casefold()
 if v in {'f','female','femenino','femenina','mujer','xx'}:return 'female'
 if v in {'m','male','masculino','hombre','xy'}:return 'male'
 return 'unknown'

def database_rows(path):
 from openpyxl import load_workbook
 path=Path(path)
 if path.is_dir():
  files=[p for p in path.glob('*.xlsx') if not p.name.startswith(('.', '~'))]
  if len(files)!=1:raise ValueError('Patient database directory must contain exactly one active .xlsx workbook')
  path=files[0]
 book=load_workbook(path,read_only=True,data_only=True);rows=[]
 try:
  for sheet in book:
   it=sheet.iter_rows(values_only=True);header=[text(v).lower() for v in next(it,())]
   if 'codigo_muestra' not in header:continue
   for values in it:
    row=dict(zip(header,values));code=text(row.get('codigo_muestra'))
    if not code:continue
    # Export pseudonymous sample codes; do not copy names, national IDs or contact details.
    age=text(row.get('edad_diagnostico'));date=row.get('fecha_diagnostico')
    rows.append(dict(sample_id=canonical(code),name=code,age=age,sex=sex(row.get('sexo')),birth_year='',
      age_context='at diagnosis' if age else '',age_as_of=date.date().isoformat() if isinstance(date,datetime.datetime) else '',source='patient database'))
 finally:book.close()
 return rows

def load_rows(path):
 with open(path,newline='') as f:
  reader=csv.DictReader(f,delimiter='\t')
  if not {'name','age'}.issubset(reader.fieldnames or []):raise ValueError('Patient TSV requires name and age columns; age may be blank. sample_id and sex are recommended.')
  return list(reader)

def write_rows(path,rows):
 path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
 with path.open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=FIELDS,delimiter='\t',extrasaction='ignore');w.writeheader();w.writerows(rows)

def prepare(c):
 records={}
 for source,rows in [('database',database_rows(c['patient_database']) if c.get('patient_database') else []),('tsv',load_rows(c['patients']) if c.get('patients') else [])]:
  seen=set()
  for row in rows:
   k=canonical(row.get('sample_id') or row.get('name'))
   if not k:continue
   if k in seen:raise ValueError('Ambiguous duplicate sample ID in patient '+source+': '+k)
   seen.add(k);records[k]={**records.get(k,{}),**{a:text(v) for a,v in row.items()}}
 key=canonical(c['sample']);r=records.get(key,{'sample_id':c['sample'],'name':c['sample'],'age':'','sex':'unknown'})
 r=dict(r);r.update(sample_id=c['sample'],sex=sex(r.get('sex')),metadata_matched=key in records)
 if r.get('birth_year'):
  year=int(r['birth_year']);current=datetime.date.today().year
  if not current-120<=year<=current:raise ValueError('Invalid birth year for '+c['sample'])
  r.update(age_min=max(0,current-year-1),age_max=current-year,age_display=f'{max(0,current-year-1)}–{current-year} years (birth year only)')
 elif text(r.get('age')):
  age=float(r['age'])
  if not 0<=age<=120:raise ValueError('Invalid age for '+c['sample'])
  r['age_display']=f'{age:g} years'+(' '+r['age_context'] if r.get('age_context') else '')
 else:r['age_display']='Not available'
 r['sex_source']='provided metadata' if r['sex']!='unknown' else 'not available'
 r['note']='A sample code/name is required for matching. Age and sex are optional; missing ages remain included. Names and ages are never inferred from sequencing.'
 c['patient_metadata']=r
 records.setdefault(key,{k:r.get(k,'') for k in FIELDS})
 output=Path(c['output'])/'pipeline_info/patients.tsv';write_rows(output,[r])
 c['patient_table_export']=str(output)

def sequencing_sex(coverage):
 # Broad QC only; no karyotype or gender identity diagnosis.
 depths={}
 if coverage and Path(coverage).is_file():
  with open(coverage) as f:
   for line in f:
    if line.startswith('#'):continue
    x=line.rstrip().split('\t')
    if len(x)>=7:depths[x[0].removeprefix('chr')]=float(x[6])
 auto=[depths[str(i)] for i in range(1,23) if depths.get(str(i),0)>0]
 if len(auto)<10 or 'X' not in depths or 'Y' not in depths:return {'status':'unresolved','sex':'unknown','reason':'Insufficient chromosome coverage statistics'}
 import statistics
 baseline=statistics.median(auto)
 if baseline<.5:return {'status':'unresolved','sex':'unknown','reason':'Autosomal depth below 0.5x; insufficient for this exploratory chromosome-coverage check'}
 xr=depths['X']/baseline;yr=depths['Y']/baseline
 estimate='female' if .75<=xr<=1.3 and yr<.1 else 'male' if .3<=xr<=.7 and yr>=.1 else 'unknown'
 return {'status':'estimated' if estimate!='unknown' else 'unresolved','sex':estimate,'x_autosome_depth_ratio':xr,'y_autosome_depth_ratio':yr,
         'reason':'Exploratory XX-like/XY-like depth pattern; not a karyotype, gender identity, or clinical sex determination. Low coverage, mapping bias and mosaicism can affect this check.'}

def resolve(c,coverage=None):
 r=dict(c.get('patient_metadata') or {'sample_id':c['sample'],'name':c['sample'],'sex':'unknown','age_display':'Not available'});r['sequencing_sex_check']=sequencing_sex(coverage)
 if r.get('sex','unknown')=='unknown' and r['sequencing_sex_check']['status']=='estimated':r['sex']=r['sequencing_sex_check']['sex'];r['sex_source']='estimated from chromosome coverage'
 return r
